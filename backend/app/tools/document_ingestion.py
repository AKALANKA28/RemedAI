from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from langsmith import traceable
from pypdf import PdfReader

from backend.app.services.logging_utils import get_app_logger

logger = get_app_logger('tool.document_ingestion')


@dataclass(slots=True)
class ExtractedDocument:
    path: str
    text: str


DATE_FORMATS = ('%Y-%m-%d', '%d-%m-%Y', '%d %B %Y', '%d %b %Y')


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding='utf-8')


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    return '\n'.join(parts)


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or '')
    return '\n'.join(pages)


def extract_text_from_path(path: str) -> ExtractedDocument:
    """Read a supported local document type and return normalized text."""
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f'Input path does not exist: {path}')
    suffix = file_path.suffix.lower()
    if suffix in {'.md', '.txt'}:
        text = _read_text_file(file_path)
    elif suffix == '.json':
        text = json.dumps(json.loads(file_path.read_text(encoding='utf-8')), indent=2)
    elif suffix == '.docx':
        text = _read_docx(file_path)
    elif suffix == '.pdf':
        text = _read_pdf(file_path)
    else:
        raise ValueError(f'Unsupported file type: {suffix}')
    return ExtractedDocument(path=str(file_path), text=text)


def _parse_date(raw: str | None) -> str | None:
    if not raw:
        return None
    cleaned = raw.strip().replace('.', '')
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _parse_budget(text: str) -> int | None:
    match = re.search(r'LKR[^0-9]*([0-9,]+)', text, flags=re.IGNORECASE)
    if not match:
        return None
    return int(match.group(1).replace(',', ''))


def _extract_section_lines(text: str, heading: str) -> list[str]:
    pattern = rf'##\s+{re.escape(heading)}\n(.*?)(?:\n##\s+|$)'
    match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return []
    block = match.group(1)
    return [line.strip('- ').strip() for line in block.splitlines() if line.strip().startswith('-')]


def _infer_requirements_from_text(text: str) -> list[dict[str, Any]]:
    candidate_lines: list[str] = []
    requirement_terms = (
        'shall',
        'must',
        'required',
        'eligible',
        'qualification',
        'experience',
        'submission',
        'bidder',
        'tenderer',
        'documents',
    )
    for raw_line in text.splitlines():
        line = ' '.join(raw_line.split())
        lowered = line.lower()
        if 35 <= len(line) <= 240 and any(term in lowered for term in requirement_terms):
            candidate_lines.append(line)
        if len(candidate_lines) >= 8:
            break

    if not candidate_lines:
        candidate_lines = ['Manual review of the official tender notice is required before bidding.']

    return [
        {
            'requirement_id': f'REQ-{index:02d}',
            'text': line,
            'mandatory': True,
            'category': 'general',
            'evidence_needed': 'Evidence or statement proving this requirement can be delivered.',
        }
        for index, line in enumerate(candidate_lines, start=1)
    ]


@traceable(name='parse_tender_package', run_type='tool')
def parse_tender_package(input_paths: list[str]) -> dict[str, Any]:
    """Extract key tender fields from a local package of markdown, text, docx, or pdf files."""
    logger.info('Parse tender package start files=%s', len(input_paths))
    documents = [extract_text_from_path(path) for path in input_paths]
    combined = '\n\n'.join(doc.text for doc in documents)
    title_match = re.search(r'\*{0,2}Project title:\*{0,2}\s*(.+)', combined, flags=re.IGNORECASE)
    issuer_match = re.search(r'\*{0,2}Issuer:\*{0,2}\s*(.+)', combined, flags=re.IGNORECASE)
    deadline_match = re.search(r'\*{0,2}Submission deadline:\*{0,2}\s*(.+)', combined, flags=re.IGNORECASE)
    duration_match = re.search(r'\*{0,2}Contract duration:\*{0,2}\s*(\d+)\s*months?', combined, flags=re.IGNORECASE)
    requirements: list[dict[str, Any]] = []
    for line in _extract_section_lines(combined, 'Mandatory Requirements'):
        req_match = re.match(r'(REQ-\d+):\s*(.+)', line)
        requirement_id = req_match.group(1) if req_match else f'REQ-{len(requirements)+1:02d}'
        text = req_match.group(2) if req_match else line
        lowered = text.lower()
        if any(token in lowered for token in ('iso', 'certification')):
            category = 'certification'
        elif 'project' in lowered or 'experience' in lowered:
            category = 'experience'
        elif any(token in lowered for token in ('mobilize', 'response', 'coverage')):
            category = 'operations'
        elif any(token in lowered for token in ('sample', 'custody', 'laboratory', 'qa')):
            category = 'quality'
        elif any(token in lowered for token in ('telemetry', 'scada', 'mapping', 'drone', 'instrumentation')):
            category = 'technical'
        else:
            category = 'general'
        requirements.append(
            {
                'requirement_id': requirement_id,
                'text': text,
                'mandatory': True,
                'category': category,
                'evidence_needed': 'Evidence or statement proving this requirement can be delivered.',
            }
        )
    if not requirements:
        requirements = _infer_requirements_from_text(combined)
    scope_lines = _extract_section_lines(combined, 'Scope of Work')
    payload = {
        'documents': [{'path': doc.path, 'chars': len(doc.text)} for doc in documents],
        'raw_text_excerpt': combined[:3500],
        'project_title': title_match.group(1).strip().strip('* ').strip() if title_match else 'Unknown project',
        'issuer': issuer_match.group(1).strip().strip('* ').strip() if issuer_match else 'Unknown issuer',
        'submission_deadline': _parse_date(deadline_match.group(1)) if deadline_match else None,
        'estimated_budget_lkr': _parse_budget(combined),
        'contract_duration_months': int(duration_match.group(1)) if duration_match else None,
        'requirements': requirements,
        'evaluation_criteria': _extract_section_lines(combined, 'Evaluation Criteria'),
        'scope_summary': scope_lines,
        'ambiguities': _extract_section_lines(combined, 'Clarifications'),
    }
    logger.info(
        'Parse tender package complete title=%s docs=%s requirements=%s chars=%s',
        payload['project_title'],
        len(documents),
        len(requirements),
        len(combined),
    )
    return payload
