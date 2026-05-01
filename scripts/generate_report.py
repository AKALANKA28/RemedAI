from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / 'docs'
REPORT_DOCX = DOCS_DIR / 'CTSE_Assignment2_Report.docx'


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for key, value in {'top': top, 'start': start, 'bottom': bottom, 'end': end}.items():
        node = tc_mar.find(qn(f'w:{key}'))
        if node is None:
            node = OxmlElement(f'w:{key}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('RemedAI')
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(17, 52, 86)
    s = doc.add_paragraph('Local Multi-Agent Tender Qualification System for Environmental Remediation Projects')
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.runs[0].italic = True
    s.runs[0].font.size = Pt(11)
    m = doc.add_paragraph('CTSE Assignment 2 Report\nReplace team member placeholders and repository URL before final submission.')
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in m.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(90, 90, 90)
    doc.add_paragraph()


def add_flow_table(doc: Document) -> None:
    doc.add_paragraph('System Workflow', style='Heading 1')
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    labels = ['Orchestrator', 'Intake', 'Compliance', 'Risk', 'Planner']
    fills = ['DCEBF7', 'E8F4EA', 'FFF1D6', 'FCE4EC', 'EDE7F6']
    for index, label in enumerate(labels):
        cell = table.rows[0].cells[index]
        cell.text = label
        shade_cell(cell, fills[index])
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.runs[0].bold = True
    doc.add_paragraph(
        'Each agent uses at least one custom Python tool, while LangGraph passes shared state and checkpoints between nodes.',
        style='Intense Quote',
    )


def add_agent_table(doc: Document) -> None:
    doc.add_paragraph('Agents and Tools', style='Heading 1')
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Agent', 'Role', 'Primary tool', 'Primary output']
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
        shade_cell(table.cell(0, idx), 'D9EAF7')
    rows = [
        ('Orchestrator', 'Registers case and decides whether to proceed', 'Case Registry Tool', 'Workflow decision'),
        ('Intake', 'Extracts tender facts and requirements', 'Document Ingestion Tool', 'Tender summary'),
        ('Compliance', 'Matches requirements against local evidence', 'Capability Lookup Tool', 'Compliance matrix'),
        ('Risk', 'Scores execution and bid risk', 'Risk Scoring Tool', 'Risk register + recommendation'),
        ('Planner', 'Builds action plan and exports artifacts', 'Plan & Export Tool', 'Submission plan + files'),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            set_cell_margins(cells[idx])
    doc.add_paragraph()


def build_report() -> Path:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    styles['Heading 1'].font.name = 'Aptos Display'
    styles['Heading 1'].font.size = Pt(15)
    styles['Heading 1'].font.color.rgb = RGBColor(17, 52, 86)
    styles['Heading 2'].font.name = 'Aptos Display'
    styles['Heading 2'].font.size = Pt(12)
    styles['Heading 2'].font.color.rgb = RGBColor(52, 86, 120)

    add_title(doc)
    doc.add_paragraph('1. Problem Domain', style='Heading 1')
    doc.add_paragraph(
        'RemedAI targets a non-generic but highly practical workflow: screening environmental remediation tenders before a bid team commits scarce proposal effort. The system is designed for firms handling hazardous-site cleanup, groundwater monitoring, and telemetry-linked field delivery. These tenders are document-heavy, schedule-sensitive, and operationally risky, which makes them a strong fit for an agentic workflow rather than a conversational assistant.'
    )
    doc.add_paragraph(
        'The chosen sample case concerns an emergency remediation and groundwater monitoring contract. The MAS must determine whether the firm can comply with mandatory criteria, identify evidence gaps, assess schedule and capacity risk, and create a clear bid/no-bid recommendation with a practical submission plan.'
    )

    add_flow_table(doc)
    add_agent_table(doc)

    doc.add_paragraph('2. Multi-Agent Architecture and Orchestration', style='Heading 1')
    doc.add_paragraph(
        'The workflow is implemented with LangGraph as a stateful directed graph. The Orchestrator Agent creates the case workspace and returns a routing decision. If the case is viable, the system proceeds sequentially through the Intake, Compliance, Risk, and Planner agents. This design intentionally keeps responsibilities distinct: extraction, evidence matching, risk interpretation, and planning are separated to reduce hallucination and make agent outputs easier to test.'
    )
    doc.add_paragraph(
        'LangChain is used for prompt construction, output parsing, and local Ollama chat-model execution. Each agent uses a small structured-output wrapper built around a local ChatOllama model. The agents are configured with strict prompts that emphasize evidence-only reasoning, explicit ambiguity handling, and concise JSON outputs suitable for downstream state handoffs.'
    )

    doc.add_paragraph('3. State Management', style='Heading 1')
    doc.add_paragraph(
        'A shared TenderMASState TypedDict stores the case id, input paths, parsed tender summary, compliance assessment, risk assessment, final plan, generated artifacts, and append-only audit trail. LangGraph reducers are used for messages and audit events, preventing accidental state loss across node transitions. The graph is compiled with a SQLite SqliteSaver checkpointer so each super-step can be recovered, inspected, and resumed if needed.'
    )

    doc.add_paragraph('4. Tooling and Observability', style='Heading 1')
    doc.add_paragraph(
        'Each agent is paired with a domain-relevant Python tool. The Intake agent reads local files and extracts structured fields from tender documents. The Compliance agent queries a local SQLite capability database. The Risk agent applies deterministic schedule, gap, and capacity rules. The Planner agent builds backward-planned tasks and writes output artifacts. The Orchestrator uses a case registry tool to persist run metadata. All tools use type hints, docstrings, and explicit error surfaces.'
    )
    doc.add_paragraph(
        'For observability, the project uses two layers. First, LangSmith tracing is enabled through environment variables, allowing each chain and tool call to appear in a trace tree. Second, a local JSONL audit log is written inside each case workspace. This dual approach ensures the observability requirement remains demonstrable even when cloud tracing is unavailable during offline marking.'
    )

    doc.add_paragraph('5. Evaluation Methodology', style='Heading 1')
    doc.add_paragraph(
        'The repository includes a unified evaluation harness implemented with pytest. Tool-level tests validate deterministic parsing and scoring logic. A graph smoke test validates the full workflow on the seeded tender. Property-based tests probe parser robustness and risk-rule stability across varying conditions. Separate per-agent tests check the most important behavioral constraints: the Intake agent must not invent requirements, the Compliance agent must surface gaps, the Risk agent must escalate unsupported mandatory requirements, and the Planner agent must produce actionable tasks and artifact files.'
    )
    doc.add_paragraph(
        'This testing strategy was chosen because local SLM outputs can vary. Instead of relying only on exact text matches, the evaluation harness checks structured invariants: required fields must exist, unsupported mandatory requirements must be escalated, and all generated plans must include named owners and due dates. These assertions make the system more reliable under model drift while still allowing natural-language flexibility.'
    )

    doc.add_paragraph('6. Deployment and Local Execution', style='Heading 1')
    doc.add_paragraph(
        'The deployment model is intentionally simple so the markers can run it on a student machine without paid infrastructure. Ollama hosts the local model, FastAPI exposes the backend workflow, and Streamlit provides a demonstration UI. Runtime artifacts, LangGraph checkpoints, and output bundles are written to local folders. Because the assignment prohibits paid provider keys, all core reasoning is performed locally with an Ollama model and local tools.'
    )
    doc.add_paragraph(
        'The recommended default model is qwen3:8b because it offers strong agent-oriented behavior and tool-use capability for a small local footprint. However, the `.env` file allows the team to switch to llama3.1:8b or another Ollama-compatible model if lab hardware constraints require it. LangSmith tracing is optional at runtime; when not configured, the local JSONL audit trail still records each agent handoff and tool output.'
    )

    doc.add_paragraph('7. Results on the Seeded Case', style='Heading 1')
    doc.add_paragraph(
        'On the included remediation tender, the system should return a conditional bid recommendation. The company has strong evidence for remediation delivery, telemetry integration, chain-of-custody controls, and mobilization speed. However, the workflow correctly identifies a recoverable gap in 24x7 response coverage and a partner-dependent gap in drone mapping support. The resulting plan prioritizes sponsor approval, evidence pack assembly, and written partner commitments before drafting the final technical proposal.'
    )

    doc.add_paragraph('8. Limitations and Future Work', style='Heading 1')
    doc.add_paragraph(
        'This version focuses on a seeded local knowledge base and a structured sample tender so the workflow is easy to reproduce. In a production extension, the system should add deeper document parsing for scanned PDFs, stronger semantic retrieval over historical evidence, and an interrupt/resume review step for procurement managers before the final recommendation is accepted.'
    )
    doc.add_paragraph(
        'Another valuable extension would be adding subcontractor-market discovery and supplier risk tracking, which would help convert conditional bids into evidence-backed partner strategies instead of static warnings. These additions were left out to keep the submission reliable, local, and easy to demonstrate within the assignment constraints.'
    )

    doc.add_paragraph('9. Team Contribution Template', style='Heading 1')
    contribution = doc.add_table(rows=1, cols=4)
    contribution.style = 'Table Grid'
    contribution.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(['Member', 'Agent owned', 'Tool owned', 'Challenges faced']):
        contribution.cell(0, idx).text = header
        shade_cell(contribution.cell(0, idx), 'E5EEF7')
    placeholders = [
        ('Student 1', 'Orchestrator', 'Case Registry Tool', 'Replace with real contribution note'),
        ('Student 2', 'Intake', 'Document Ingestion Tool', 'Replace with real contribution note'),
        ('Student 3', 'Compliance', 'Capability Lookup Tool', 'Replace with real contribution note'),
        ('Student 4', 'Risk', 'Risk Scoring Tool', 'Replace with real contribution note'),
    ]
    for row in placeholders:
        cells = contribution.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            set_cell_margins(cells[idx])

    doc.add_paragraph('Repository link: REPLACE_WITH_GITHUB_OR_GITLAB_URL', style='Intense Quote')
    doc.save(REPORT_DOCX)
    return REPORT_DOCX


if __name__ == '__main__':
    path = build_report()
    print(path)
